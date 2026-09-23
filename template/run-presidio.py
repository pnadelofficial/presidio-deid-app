from io import StringIO, BytesIO
import json
import os
import time

import streamlit as st

from docx import Document
from PIL import Image

from presidio_anonymizer.entities import OperatorConfig
from presidio_utils import create_analyzer, create_anonymizer, choose_model, create_redactor, create_dicom_redactor

st.title("Data Anonymizer")

data_type = st.selectbox("What kind of data do you want to de-identify", ["Texts", "Images"])

if data_type == "Texts":
    entity_mapping = {"PER":"PERSON"}
    string_data = None

    uploaded_file = st.file_uploader(
        "Choose a file",
    )
    if uploaded_file:
        if uploaded_file.name.endswith(".docx"):
            document = Document(BytesIO(uploaded_file.getvalue()))
            string_data = "\n".join([para.text for para in document.paragraphs])
        elif uploaded_file.name.endswith(".txt"):
            stringio = StringIO(uploaded_file.getvalue().decode("utf-8"))
            string_data = stringio.read()
        else:
            st.warning("This application only support .docx and .txt texts.")

    if string_data:
        with st.spinner("Anonymizing"):
            time.sleep(1)
            lang_code, model_name = choose_model(string_data)
            analyzer = create_analyzer(
                lang_code,
                model_name,
                entity_mapping=entity_mapping
            )
            anonymizer = create_anonymizer()

            analyzer_results = analyzer.analyze(
                text=string_data,
                language=lang_code
            )
            anonymized_results = anonymizer.anonymize(
                text=string_data,
                analyzer_results=analyzer_results,
                operators={
                    "DEFAULT": OperatorConfig("replace", {"new_value": "<ANONYMIZED>"})
                }
            )

        anon = json.loads(anonymized_results.to_json())['text']

        base_path, extension = os.path.splitext(uploaded_file.name)
        if extension == ".docx":
            anon_doc = Document()
            for para in anon.split("\n"):
                anon_doc.add_paragraph(para)
            anon_doc.save(f"{os.getcwd()}/{base_path}_deid{extension}")
        elif extension == ".txt":
            with open(f"{os.getcwd()}/{base_path}_deid{extension}", "w") as file:
                file.write(anon)


        with open(f"{os.getcwd()}/{base_path}_deid{extension}", "r" if extension == ".txt" else "rb") as file:
            st.download_button(
                label="Download De-identified data",
                data=file,
                file_name=f"{base_path}_deid{extension}"
            )

elif data_type == "Images":
    st.subheader("Image De-identification")
    st.info("📝 **Note**: This tool uses OCR to detect and redact text PHI embedded in images. For DICOM files, it will also use DICOM metadata to improve detection accuracy.")
    
    # Sub-option for image type
    image_type = st.radio(
        "Select image type",
        ["JPEG/PNG/BMP", "DICOM (.dcm)"],
        help="DICOM (Digital Imaging and Communications in Medicine) files are standard medical imaging format.",
        index=1  # Default to DICOM since that's the primary use case
    )
    
    uploaded_file = st.file_uploader(
        "Choose a file",
        type=["jpg", "jpeg", "png", "bmp", "dcm"] if image_type == "JPEG/PNG/BMP" else ["dcm"],
    )
    
    if uploaded_file:
        base_name, extension = os.path.splitext(uploaded_file.name)
        
        if image_type == "JPEG/PNG/BMP":
            # Standard image redaction
            image = Image.open(BytesIO(uploaded_file.getvalue()))
            redactor = create_redactor()
            redacted = redactor.redact(image=image)
            st.image(redacted, caption="Redacted Image")
            
            # Save and download
            # Convert RGBA to RGB if needed (JPEG doesn't support alpha channel)
            if redacted.mode == "RGBA":
                redacted = redacted.convert("RGB")
            
            output_path = f"{os.getcwd()}/{base_name}_deid.jpg"
            redacted.save(output_path)
            with open(output_path, "rb") as file:
                st.download_button(
                    label="Download Redacted Image",
                    data=file,
                    file_name=f"{base_name}_deid.jpg",
                    mime="image/jpeg"
                )
        
        elif image_type == "DICOM (.dcm)":
            # DICOM image redaction using DicomImageRedactorEngine
            try:
                import pydicom
                dicom_redactor = create_dicom_redactor()
                
                if dicom_redactor is None:
                    st.error("Failed to initialize DICOM redactor. Please ensure pydicom and presidio_image_redactor are installed.")
                else:
                    # Load DICOM file
                    dicom_instance = pydicom.dcmread(BytesIO(uploaded_file.getvalue()))
                    
                    # Redact PHI from DICOM image (fill options: "contrast", "background")
                    fill_method = st.selectbox(
                        "Fill method for redacted areas",
                        ["contrast", "background"],
                        help="'contrast' fills with black/white contrast; 'background' blends with image background."
                    )
                    
                    with st.spinner("Redacting DICOM image..."):
                        redacted_dicom = dicom_redactor.redact(dicom_instance, fill=fill_method)
                    
                    # Display original and redacted side by side
                    import matplotlib.pyplot as plt
                    fig, ax = plt.subplots(1, 2, figsize=(12, 6))
                    ax[0].imshow(dicom_instance.pixel_array, cmap="gray")
                    ax[0].set_title('Original DICOM')
                    ax[0].axis('off')
                    ax[1].imshow(redacted_dicom.pixel_array, cmap="gray")
                    ax[1].set_title(f'Redacted DICOM (fill={fill_method})')
                    ax[1].axis('off')
                    plt.tight_layout()
                    st.pyplot(fig)
                    
                    # Save and download redacted DICOM
                    output_path = f"{os.getcwd()}/{base_name}_deid.dcm"
                    pydicom.dcmwrite(output_path, redacted_dicom)
                    with open(output_path, "rb") as file:
                        st.download_button(
                            label="Download Redacted DICOM",
                            data=file,
                            file_name=f"{base_name}_deid.dcm",
                            mime="application/dicom"
                        )
                        
            except ImportError as e:
                st.error(f"Missing dependency: {e}. Please ensure pydicom is installed.")
            except Exception as e:
                st.error(f"Error processing DICOM image: {e}")